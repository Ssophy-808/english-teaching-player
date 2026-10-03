from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "exports" / "Book2-Review1-3-Day4-Worksheet.docx"


def set_font(run, name="Segoe UI", size=11, bold=False, color="111111"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def add_text(paragraph, value, size=11, bold=False, color="111111"):
    run = paragraph.add_run(value)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_bottom_border(paragraph, dashed=False):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "dashed" if dashed else "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:color"), "9FC4E7")
    pbdr.append(bottom)


def remove_paragraph_borders(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is not None:
        ppr.remove(pbdr)


def four_lines(container, groups=1):
    for _ in range(groups):
        for index in range(4):
            p = container.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.65
            add_text(p, " ", size=6)
            add_bottom_border(p, dashed=index in (1, 2))


def page_header(doc, page_title, page_number):
    p = doc.add_paragraph()
    if page_number > 1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(0)
    add_text(p, " ", size=2)
    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(title, page_title, size=19, bold=True, color="000000")
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(name, f"Name: __________________________    Page {page_number} of 4", size=10, bold=True)


def picture_prompt(cell, image=None, label=None, width=2.2):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if label:
        add_text(p, f"{label}  ", size=16, bold=True, color="17325C")
    if image:
        p.add_run().add_picture(str(ROOT / image), width=Cm(width))


def clear_cell_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tcpr.append(borders)


def question_row(doc, question, image=None, label=None, image_width=2.2):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(14.5)
    table.columns[1].width = Cm(3.0)
    left, right = table.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    add_text(p, question, size=11, bold=True)
    four_lines(left)
    picture_prompt(right, image, label, image_width)
    for cell in (left, right):
        clear_cell_borders(cell)


def compact_visual_row(doc, left_text, right_text, image, label=None):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (8.1, 2.6, 7.0)
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        clear_cell_borders(cell)
    add_text(table.cell(0, 0).paragraphs[0], left_text, size=10.5)
    picture_prompt(table.cell(0, 1), image, label, 1.5)
    add_text(table.cell(0, 2).paragraphs[0], right_text, size=10.5)


def choice_visual_row(doc, answer, choices, image, label=None, index=1):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Cm(15.3)
    right.width = Cm(2.4)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        clear_cell_borders(cell)
    p = left.paragraphs[0]
    add_text(p, f"{index}. {answer}", size=10.5, bold=True)
    p = left.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    add_text(p, choices, size=9.5)
    picture_prompt(right, image, label, 1.7)


def story_page(doc):
    page_header(doc, "Lumi and Ludi's Review Story", 1)
    p = doc.add_paragraph()
    add_text(p, "Follow the whole story. Answer every question, then read the passage aloud.", size=10, bold=True, color="17325C")
    for line in ("Hi, I'm Lumi.", "Today is the final review challenge.", "Look at the pictures and answer the questions."):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_text(p, line, size=11)
    question_row(doc, "Is there a table?", "assets/images/book2/unit1/table.jpeg")
    question_row(doc, "Are there eleven windows?", "assets/images/book2/unit1/window.jpeg", "11")
    question_row(doc, "How many yo-yos are there?", "assets/images/book2/unit3/yo-yo.jpeg", "8")
    p = doc.add_paragraph()
    add_text(p, "Check every complete sentence. Great job today!", size=11)


def matching_page(doc):
    page_header(doc, "Cumulative Dialogue", 2)
    p = doc.add_paragraph()
    add_text(p, "Match each question to the best complete answer.", size=10, bold=True, color="17325C")
    pairs = [
        ("1. Is there a door?", "A. Yes, there are.", "assets/images/book2/unit1/door.jpeg", None),
        ("2. Are there twelve books?", "B. Yes, there is. There is a fan!", "assets/images/book1/unit5/book.png", "12"),
        ("3. Is there a fan?", "C. Yes, there are.", "assets/images/book2/unit1/fan.jpeg", None),
        ("4. Are there twenty books?", "D. Yes, there is.", "assets/images/book1/unit5/book.png", "20"),
    ]
    for left, right, image, label in pairs:
        compact_visual_row(doc, left, right, image, label)
    p = doc.add_paragraph()
    add_text(p, "Choose one pair. Write the complete answer.", size=11, bold=True)
    four_lines(doc)


def choice_page(doc):
    page_header(doc, "Final Cumulative Check", 3)
    p = doc.add_paragraph()
    add_text(p, "Circle the only correct question.", size=10, bold=True, color="17325C")
    items = [
        ("Answer: No, there is not.", "A Is there a blackboard?   B Is there a blackboard   C Is there blackboard?", "assets/images/book2/unit1/blackboard.jpeg", "NO"),
        ("Answer: Yes, there are.", "A Are there seventeen books?   B Is there seventeen books?   C There are seventeen books?", "assets/images/book1/unit5/book.png", "17"),
        ("Answer: There are ten board games.", "A How many board games are there?   B How much board games?   C What board games there are?", "assets/images/book2/unit3/board-game.jpeg", "10"),
        ("Answer: No, there is not.", "A Is there a television?   B Are there a television?   C There is a television?", "assets/images/book2/unit1/television.jpeg", "NO"),
    ]
    for index, (answer, choices, image, label) in enumerate(items, 1):
        choice_visual_row(doc, answer, choices, image, label, index)
    p = doc.add_paragraph()
    add_text(p, "Choose one answer. Write the complete sentence.", size=11, bold=True)
    four_lines(doc)


def chain_page(doc):
    page_header(doc, "Independent Chain Challenge", 4)
    p = doc.add_paragraph()
    add_text(p, "Answer in complete sentences. Then write one new question that keeps the conversation going.", size=10, bold=True, color="17325C")
    questions = [
        ("Is there a whiteboard?", "assets/images/book2/unit1/whiteboard.jpeg", None),
        ("Are there fourteen books?", "assets/images/book1/unit5/book.png", "14"),
        ("How many video games are there?", "assets/images/book2/unit3/video-game.png", "9"),
        ("Is there a telephone?", "assets/images/book2/unit1/telephone.jpeg", None),
        ("Are there nineteen windows?", "assets/images/book2/unit1/window.jpeg", "19"),
        ("How many blocks are there?", "assets/images/book2/unit3/blocks.jpeg", "16"),
    ]
    for index, (question, image, label) in enumerate(questions, 1):
        question_row(doc, f"{index}. {question}", image, label, 1.8)
    p = doc.add_paragraph()
    add_text(p, "My next question:", size=11, bold=True)
    four_lines(doc)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(11)
    title_style_ppr = doc.styles["Title"].element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)

    story_page(doc)
    matching_page(doc)
    choice_page(doc)
    chain_page(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
